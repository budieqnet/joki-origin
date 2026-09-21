import csv
import os

_DOCX_MISSING = _XLSX_MISSING = _PPTX_MISSING = _PDF_MISSING = False


def _check_docx():
    global _DOCX_MISSING
    try:
        import docx
        return docx
    except ImportError:
        _DOCX_MISSING = True
        return None


def _check_xlsx():
    global _XLSX_MISSING
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        _XLSX_MISSING = True
        return None


def _check_pptx():
    global _PPTX_MISSING
    try:
        from pptx import Presentation
        return Presentation
    except ImportError:
        _PPTX_MISSING = True
        return None


def _check_pdfplumber():
    global _PDF_MISSING
    try:
        import pdfplumber
        return pdfplumber
    except ImportError:
        _PDF_MISSING = True
        return None


def _check_fitz():
    global _PDF_MISSING
    try:
        import fitz
        return fitz
    except ImportError:
        _PDF_MISSING = True
        return None


def _missing_msg():
    parts = []
    if _DOCX_MISSING:
        parts.append("DOCX: pip install python-docx")
    if _XLSX_MISSING:
        parts.append("XLSX: pip install openpyxl")
    if _PPTX_MISSING:
        parts.append("PPTX: pip install python-pptx")
    if _PDF_MISSING:
        parts.append("PDF: pip install pdfplumber PyMuPDF")
    if parts:
        return "Library belum terinstall. Install dulu:\n" + "\n".join(parts)
    return ""


def handle_read_office(args):
    path = args.get("path", "")
    if not os.path.isfile(path):
        return f"Error: File tidak ditemukan: {path}"

    ext = os.path.splitext(path)[1].lower()

    if ext == ".docx":
        return _read_docx(path)
    elif ext == ".xlsx":
        return _read_xlsx(path)
    elif ext == ".pptx":
        return _read_pptx(path)
    elif ext == ".pdf":
        return _read_pdf(path)
    elif ext == ".csv":
        return _read_csv(path)
    else:
        return f"Error: Format '{ext}' tidak didukung. Supported: .docx, .xlsx, .pptx, .pdf, .csv"


def _read_docx(path):
    docx = _check_docx()
    if not docx:
        return _missing_msg()
    doc = docx.Document(path)
    lines = []
    for para in doc.paragraphs:
        lines.append(para.text)
    for table in doc.tables:
        lines.append("")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(cells))
        lines.append("")
    text = "\n".join(lines).strip()
    if not text:
        return "(dokumen kosong)"
    info = f"File: {path}\nFormat: DOCX\n"
    info += f"Total paragraf: {len(doc.paragraphs)}"
    if doc.tables:
        info += f", Tabel: {len(doc.tables)}"
    info += f"\n\n{text}"
    return info


def _read_xlsx(path):
    openpyxl = _check_xlsx()
    if not openpyxl:
        return _missing_msg()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = []
    sheets = wb.sheetnames
    for si, sn in enumerate(sheets):
        if si > 0:
            lines.append("")
        lines.append(f"=== Sheet: {sn} ===")
        ws = wb[sn]
        if ws.max_row is None or ws.max_column is None:
            lines.append("(kosong)")
            continue
        for row in ws.iter_rows(values_only=True):
            vals = [
                str(v).strip() if v is not None else ""
                for v in row
            ]
            line = " | ".join(vals)
            if line.strip():
                lines.append(line)
    wb.close()
    text = "\n".join(lines).strip()
    if not text:
        return "(file kosong)"
    info = f"File: {path}\nFormat: XLSX\nSheets: {', '.join(sheets)}\n\n{text}"
    return info


def _read_pptx(path):
    Presentation = _check_pptx()
    if not Presentation:
        return _missing_msg()
    prs = Presentation(path)
    lines = []
    for si, slide in enumerate(prs.slides, 1):
        lines.append(f"\n--- Slide {si} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:  # type: ignore[union-attr]
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        lines.append(t)
            if shape.has_table:  # type: ignore[union-attr]
                table = shape.table  # type: ignore[union-attr]
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    lines.append(" | ".join(cells))
    text = "\n".join(lines).strip()
    if not text:
        return "(presentasi kosong)"
    info = f"File: {path}\nFormat: PPTX\nTotal slide: {len(prs.slides)}\n\n{text}"
    return info


def _read_pdf(path):
    text = _read_pdf_pdfplumber(path)
    if text is None:
        text = _read_pdf_fitz(path)
    if text is None:
        return _missing_msg()
    info = f"File: {path}\nFormat: PDF\n\n{text}"
    return info


def _read_pdf_pdfplumber(path):
    pdfplumber = _check_pdfplumber()
    if not pdfplumber:
        return None
    try:
        lines = []
        with pdfplumber.open(path) as pdf:
            for pi, page in enumerate(pdf.pages, 1):
                t = page.extract_text()
                if t and t.strip():
                    lines.append(f"--- Halaman {pi} ---")
                    lines.append(t.strip())
        text = "\n".join(lines).strip()
        return text if text else "(PDF kosong)"
    except Exception:  # noqa: BLE001
        return None


def _read_pdf_fitz(path):
    fitz = _check_fitz()
    if not fitz:
        return None
    try:
        lines = []
        doc = fitz.open(path)
        for pi in range(len(doc)):
            page = doc[pi]
            t = page.get_text()
            if t and t.strip():
                lines.append(f"--- Halaman {pi + 1} ---")
                lines.append(t.strip())
        doc.close()
        text = "\n".join(lines).strip()
        return text if text else "(PDF kosong)"
    except Exception:  # noqa: BLE001
        return None


def handle_write_office(args):
    path = args.get("path", "")
    if not path:
        return "Error: Parameter 'path' wajib diisi."
    ext = os.path.splitext(path)[1].lower()
    content = args.get("content", "")
    if ext == ".docx":
        return _write_docx(path, content, args)
    elif ext == ".xlsx":
        return _write_xlsx(path, content, args)
    elif ext == ".pptx":
        return _write_pptx(path, content, args)
    elif ext == ".csv":
        return _write_csv(path, content, args)
    else:
        return f"Error: Format '{ext}' tidak didukung untuk write. Supported: .docx, .xlsx, .pptx, .csv"


def _write_docx(path, content, args):
    docx = _check_docx()
    if not docx:
        return _missing_msg()
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    doc = docx.Document()
    for para_text in content.split("\n"):
        doc.add_paragraph(para_text)
    tables_raw = args.get("tables", "")
    if tables_raw:
        import json
        try:
            tables = json.loads(tables_raw) if isinstance(tables_raw, str) else tables_raw
            for table_data in tables:
                if not table_data:
                    continue
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                for ri, row_data in enumerate(table_data):
                    for ci, cell_val in enumerate(row_data):
                        table.rows[ri].cells[ci].text = str(cell_val)
        except (json.JSONDecodeError, IndexError):
            pass
    doc.save(path)
    sz = os.path.getsize(path)
    return f"DOCX created: {path} ({sz} bytes, {len(content.splitlines())} paragraf)"


def _write_xlsx(path, content, args):
    openpyxl = _check_xlsx()
    if not openpyxl:
        return _missing_msg()
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb = openpyxl.Workbook()
    sheet_name = args.get("sheet_name", "Sheet1")
    ws = wb.active
    ws.title = sheet_name
    rows_args = args.get("rows", "")
    if rows_args:
        import json
        try:
            all_rows = json.loads(rows_args) if isinstance(rows_args, str) else rows_args
            for ri, row_data in enumerate(all_rows, 1):
                for ci, cell_val in enumerate(row_data, 1):
                    ws.cell(row=ri, column=ci, value=cell_val)
        except (json.JSONDecodeError, TypeError):
            pass
    elif content:
        for ri, line in enumerate(content.split("\n"), 1):
            cells = [c.strip() for c in line.replace("\t", ",").split(",")]
            for ci, cell_val in enumerate(cells, 1):
                ws.cell(row=ri, column=ci, value=cell_val)
    wb.save(path)
    sz = os.path.getsize(path)
    return f"XLSX created: {path} ({sz} bytes)"


def _write_pptx(path, content, args):
    Presentation = _check_pptx()
    if not Presentation:
        return _missing_msg()
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    prs = Presentation()
    title_text = args.get("title", "Presentation")
    slide_blocks = [s.strip() for s in content.split("---")] if content else [""]
    for si, block in enumerate(slide_blocks):
        if si == 0:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            title = slide.shapes.title  # type: ignore[union-attr]
            if title:
                title.text = title_text  # type: ignore[union-attr]
            subtitle = slide.placeholders[1] if len(slide.placeholders) > 1 else None
            if subtitle:
                subtitle.text = block if block else "Created by Joki"  # type: ignore[union-attr]
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            title = slide.shapes.title  # type: ignore[union-attr]
            if title:
                title.text = f"Slide {si}"  # type: ignore[union-attr]
            content_box = slide.placeholders[1] if len(slide.placeholders) > 1 else None
            if content_box and block:
                content_box.text = block  # type: ignore[union-attr]
    prs.save(path)
    sz = os.path.getsize(path)
    num_slides = len(slide_blocks)
    return f"PPTX created: {path} ({sz} bytes, {num_slides} slide)"


def _write_csv(path, content, wr_args):
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    rows_args = wr_args.get("rows", "")
    if rows_args:
        import json
        try:
            all_rows = json.loads(rows_args) if isinstance(rows_args, str) else rows_args
            with open(path, "w", newline="") as f:
                writer = csv.writer(f)
                writer.writerows(all_rows)
            sz = os.path.getsize(path)
            return f"CSV created: {path} ({sz} bytes, {len(all_rows)} baris)"
        except (json.JSONDecodeError, TypeError):
            pass
    with open(path, "w", newline="") as f:
        f.write(content)
    sz = os.path.getsize(path)
    return f"CSV created: {path} ({sz} bytes, {len(content.splitlines())} baris)"


def _read_csv(path):
    try:
        with open(path, newline="", encoding="utf-8") as f:
            sample = f.read(4096)
            f.seek(0)
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            reader = csv.reader(f, dialect)
            lines = []
            for row in reader:
                lines.append(" | ".join(row))
            text = "\n".join(lines).strip()
    except Exception:  # noqa: BLE001
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            lines = []
            for row in reader:
                lines.append(" | ".join(row))
            text = "\n".join(lines).strip()
    if not text:
        return "(CSV kosong)"
    info = f"File: {path}\nFormat: CSV\nBaris: {len(lines)}\n\n{text}"
    return info
